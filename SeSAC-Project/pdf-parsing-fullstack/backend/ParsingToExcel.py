# /// script
# requires-python = ">=3.10"
# dependencies = [
#      "pandas",
#      "openpyxl",
#      "pymupdf",
#      "python-docx",
#      "pyhwp",
#      "python-dotenv",
# ]
# ///

import os
import re
import subprocess
import glob
import json
import urllib.request
import urllib.error
import fitz
import pandas as pd

try:
    from dotenv import load_dotenv
except ImportError:
    load_dotenv = None

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILENAME = "지원자_통합관리.xlsx"
OUTPUT_PATH = os.path.join(BASE_DIR, OUTPUT_FILENAME)
if load_dotenv:
    load_dotenv(os.path.join(BASE_DIR, ".env"))

# ─── 공통 정규식 ────────────────────────────────────────────────────────────
# 날짜 범위: 2020.03~현재 / 2020.03-2024.01 / 2020.03 ~ 2024.01 등
_DATE_RANGE = re.compile(
    r"(\d{4}[.\-/]\d{1,2}"
    r"\s*[~\-–—]\s*"
    r"(?:\d{4}[.\-/]\d{1,2}|현재|재직|재학|중|진행))"
)
# 단독 연월: 2020.03 / 2020-03
_YEARMONTH = re.compile(r"\d{4}[.\-]\d{1,2}")
# 학교명
_SCHOOL = re.compile(r"([가-힣]{2,}\s*(?:대학교|대학원|대학))")
# 전공
_MAJOR_PATS = [
    re.compile(r"([가-힣]{2,}\s*(?:학과|학부|공학과|공학부|전공|계열))"),
    re.compile(r"([가-힣]{2,}(?:공학|과학|시스템|전자|기계|경영|경제|컴퓨터|소프트웨어|정보|통신|디자인|미디어|통계))"),
]
_GRAD_STATUS = ["졸업", "재학", "수료", "중퇴", "편입", "예정"]
_EDU_SKIP = {"학력사항", "학력구분", "학교명", "전공", "재학기간", "소재지", "학점",
             "졸업구분", "입학", "구분", "기간", "학교", "학과명", "전공 및 학위구분", "-",
             "단과대학", "단과대학(학부)", "학부", "학 과", "평균학점", "연봉", "직급",
             "부서", "담당업무", "회사명", "기 간", "년", "월", "성명", "생년월일"}
_CAREER_KWS = [
    "MCU", "펌웨어", "개발", "사원", "대리", "과장", "차장", "부장", "팀장",
    "엔지니어", "Engineer", "연구", "설계", "제어", "테크", "솔루션",
    "오토", "센서", "R&D", "시스템", "소프트웨어", "하드웨어", "매니저",
    "직원", "재직", "근무", "입사", "퇴사", "(주)", "㈜", "주식회사",
    "회사", "기업", "산업", "인턴", "계약직", "정규직", "선임", "수석", "책임",
]
_EDU_CTX_KWS = ["대학교", "대학원", "대학", "고등학교", "수료", "양성과정", "사관학교", "중학교"]
_NOISE_LINE_KWS = {
    "학력사항", "학력구분", "학교명", "단과대학", "단과대학(학부)", "학 과", "평균학점", "소재지",
    "경력사항", "회사명", "직 급", "부 서", "담당업무", "연 봉", "해외경험", "외국어",
    "자격사항", "병역사항", "자기소개", "기 간", "기간", "년", "월", "구분", "평가년월",
    "활 동  내 역", "자격종류", "등급", "기 관", "국가보훈", "대상여부",
}


def _parse_edu_from_text(text: str, all_dates: list[str]) -> tuple[str, str]:
    """텍스트에서 최종학력과 학력날짜 추출 (표 없는 포맷 대응 폴백)"""
    best_edu = ""
    best_date = ""
    best_score = -1
    # 학교명을 손상시키지 않는 멀티워드 헤더만 제거 ('학교' 단독 단어는 제외)
    _safe_skip = ["학력사항", "학력구분", "학교명", "재학기간", "졸업구분", "학점",
                  "소재지", "학력사항", "전공 및 학위구분"]
    lines = text.split("\n")

    for i, line in enumerate(lines):
        if not any(kw in line for kw in ["대학교", "대학원", "대학"]):
            continue
        temp = line
        for h in _safe_skip:
            temp = temp.replace(h, "")
        m = _SCHOOL.search(temp)
        if not m:
            continue
        school_name = m.group(1).replace(" ", "")
        if school_name in _EDU_SKIP or "단과대학" in school_name:
            continue
        ctx = " ".join(lines[max(0, i - 2): i + 4])

        major_name = ""
        remaining = temp[m.end():] or ctx
        for mp in _MAJOR_PATS:
            mm = mp.search(remaining)
            if mm:
                major_name = mm.group(1).strip()
                break

        status = next((gs for gs in _GRAD_STATUS if gs in ctx), "")

        edu_date = ""
        school_pos = text.find(school_name)
        best_dist = 10**9
        for d in all_dates:
            d_pos = text.find(d)
            if d_pos < 0:
                continue
            if d in ctx or abs(d_pos - school_pos) < 250:
                dist = abs(d_pos - school_pos)
                if dist < best_dist:
                    best_dist = dist
                    edu_date = d

        # 최종학력은 "학교명 + 기간"이 함께 있는 항목만 채택
        if not edu_date:
            continue

        score = 0
        if "대학원" in school_name:
            score += 1
        elif "대학교" in school_name or school_name.endswith("대학"):
            score += 2
        if major_name:
            score += 1
        if status:
            score += 1

        candidate = _build_edu_str(school_name, major_name, status, edu_date)
        if score > best_score:
            best_score = score
            best_edu = candidate
            best_date = edu_date
    return best_edu, best_date


def _parse_career_from_text(
    text: str, all_dates: list[str], edu_date: str
) -> tuple[str, str, str]:
    """텍스트에서 경력기간/회사/직무 추출 (표 없는 포맷 폴백)
    반환: (경력기간, 경력회사, 경력직무)  각각 \\n 구분
    """
    entries: list[dict] = []
    seen: set[str] = set()

    for d in all_dates:
        if d == edu_date or d in seen:
            continue
        pos = text.find(d)
        ctx = text[max(0, pos - 200): pos + 200]

        if not any(ckw.upper() in ctx.upper() for ckw in _CAREER_KWS):
            continue
        if any(ekw in ctx for ekw in _EDU_CTX_KWS):
            continue

        seen.add(d)
        company, job = "", ""

        # 날짜가 포함된 줄에서 회사명/직무 추출
        for line in ctx.split("\n"):
            if d not in line:
                continue
            line = line.strip()
            date_pos = line.find(d)
            before = line[:date_pos].strip()
            after = line[date_pos + len(d):].strip()
            combined = (before + " " + after).strip()

            # 회사명: 첫 한글 단어 그룹 (일반적으로 줄 앞쪽에 위치)
            comp_m = re.search(r"([가-힣]{2,}[\w가-힣]*)", combined)
            if comp_m:
                company = comp_m.group(1)

            # 직무: 직무 관련 키워드를 포함하는 한글 구문
            job_m = re.search(
                r"([가-힣/]{2,}(?:개발|분석|관리|기획|설계|마케팅|영업|연구|사원|대리|과장|팀장|엔지니어))",
                combined,
            )
            if job_m:
                job = job_m.group(1)
            break

        entries.append({"period": d, "company": company, "job": job})

    return (
        "\n".join(e["period"]  for e in entries),
        "\n".join(e["company"] for e in entries),
        "\n".join(e["job"]     for e in entries),
    )


def _extract_basic(text: str) -> dict:
    """텍스트에서 기본 정보 추출 (이름/생년/연락처/이메일/주소/지원직무)"""
    name_m = re.search(
        r"(?:성\s*명|이\s*름|이름|Name)\s*(?:\([^)]*\))?\s*[:：]?\s*([가-힣]{2,5})", text
    )
    birth_m = re.search(r"(\d{4}[.\-/]\d{1,2}[.\-/]\d{1,2})", text)
    phone_m = re.search(
        r"((?:010|011|016|017|018|019)[-\s]?\d{3,4}[-\s]?\d{4})",
        text,
    ) or re.search(
        r"((?:02|031|032|033|041|042|043|044|051|052|053|054|055|061|062|063|064|070)"
        r"[-\s]?\d{3,4}[-\s]?\d{4})",
        text,
    )
    email_m = re.search(r"([a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,})", text)
    addr_m = re.search(r"(?:주\s*소|거\s*주\s*지|Address)\s*[:：]?\s*([^\n]{5,})", text)
    job_m = re.search(
        r"(?:지\s*원\s*직\s*무|지\s*원\s*분\s*야|직\s*무|Position|지원직무)\s*[:：]?\s*"
        r"([가-힣\s\w/.\(\)]+?)(?=\n|경력여부|희망연봉|희망|About|성명|생년|$)",
        text,
    )
    return {
        "이름": name_m.group(1).strip() if name_m else "",
        "생년월일": birth_m.group(1) if birth_m else "",
        "연락처": phone_m.group(1).strip() if phone_m else "",
        "이메일": email_m.group(1) if email_m else "",
        "주소": addr_m.group(1).strip() if addr_m else "",
        "원문 지원직무": job_m.group(1).strip() if job_m else "",
    }


def _normalize_hwp_table_text(text: str) -> str:
    """HWP 표가 줄 단위로 분해된 경우를 날짜 중심으로 정규화"""
    normalized = text
    # 이메일 분절: 한글/영문 로컬파트 모두 co\nm -> com 복원
    normalized = re.sub(
        r"([a-zA-Z0-9._%+\-가-힣]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2})\s*\n\s*([a-zA-Z]{1,3})",
        r"\1\2",
        normalized,
    )
    # 2026 \n 04 -> 2026.04
    normalized = re.sub(r"(\d{4})\s*\n\s*(\d{1,2})(?=\s*(?:\n|$))", r"\1.\2", normalized)
    # 2026.04 \n - \n 2026.07 -> 2026.04~2026.07
    normalized = re.sub(
        r"(\d{4}[.\-]\d{1,2})\s*\n\s*[-~]\s*\n\s*(\d{4}[.\-]\d{1,2})",
        r"\1~\2",
        normalized,
    )
    return normalized


def _normalize_pdf_text(text: str) -> str:
    """PDF 줄바꿈 분절 텍스트를 파싱 친화 형태로 정규화"""
    normalized = text
    # 이메일 분절: 한글/영문 로컬파트 모두 co\nm -> com 복원
    normalized = re.sub(
        r"([a-zA-Z0-9._%+\-가-힣]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2})\s*\n\s*([a-zA-Z]{1,3})",
        r"\1\2",
        normalized,
    )
    # 휴대폰 분절: 010-1234-1\n234 -> 010-1234-1234
    normalized = re.sub(
        r"((?:010|011|016|017|018|019)[-\s]?\d{3,4}[-\s]?\d{1,2})\s*\n\s*(\d{2,4})",
        r"\1\2",
        normalized,
    )
    # 연월 분절: 2026 04 -> 2026.04
    normalized = re.sub(r"(\d{4})\s+(\d{1,2})(?=\s*(?:\n|$))", r"\1.\2", normalized)
    # 기간 분절: YYYY.MM \n - \n YYYY.MM -> YYYY.MM~YYYY.MM
    normalized = re.sub(
        r"(\d{4}[.\-]\d{1,2})\s*\n\s*[-~]\s*\n\s*(\d{4}[.\-]\d{1,2})",
        r"\1~\2",
        normalized,
    )
    # 같은 줄 변형: YYYY.MM - YYYY.MM -> YYYY.MM~YYYY.MM
    normalized = re.sub(
        r"(\d{4}[.\-]\d{1,2})\s*[-~]\s*(\d{4}[.\-]\d{1,2})",
        r"\1~\2",
        normalized,
    )
    return normalized


def _parse_hwp_sections(text: str) -> tuple[str, str, str, str]:
    """HWP 표형/자유형 이력서에서 학력/경력을 줄 단위로 보강 추출"""
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
    if not lines:
        return "", "", "", ""

    def _find_index(keyword: str) -> int:
        for i, ln in enumerate(lines):
            if keyword in ln:
                return i
        return -1

    edu_start = _find_index("학력사항")
    career_start = _find_index("경력사항")
    career_end = _find_index("해외경험")

    def _line_has_header(line: str) -> bool:
        return any(h in line for h in _NOISE_LINE_KWS)

    # 학력: 마지막 대학(원) 행 + 인접 날짜 조합
    edu_str = ""
    edu_date = ""
    best_edu_score = -1
    edu_scan_lines = (
        lines[edu_start + 1: career_start if career_start > edu_start else len(lines)]
        if edu_start >= 0 else lines
    )
    for i, ln in enumerate(edu_scan_lines):
        if not re.fullmatch(r"[가-힣A-Za-z0-9() ]{2,}\s*(?:대학교|대학원|대학)", ln):
            continue
        if ln in _EDU_SKIP or "단과대학" in ln:
            continue

        major = ""
        for nxt in edu_scan_lines[i + 1:i + 5]:
            if _line_has_header(nxt):
                continue
            if any(tag in nxt for tag in ["학과", "학부", "전공", "계열"]):
                major = nxt
                break

        # 주변 10줄에서 y,m,y,m 패턴 찾기
        window = " ".join(edu_scan_lines[max(0, i - 10): i + 2])
        nums = re.findall(r"\b(\d{4})[.\-]?(\d{1,2})\b", window)
        period = ""
        if len(nums) >= 2:
            s_y, s_m = nums[-2]
            e_y, e_m = nums[-1]
            period = f"{s_y}.{s_m}~{e_y}.{e_m}"

        status = next((gs for gs in _GRAD_STATUS if gs in " ".join(edu_scan_lines[i:i + 8])), "")
        # "학교명 + 기간"이 있어야 최종학력 후보로 인정
        if not period:
            continue

        school = ln.replace(" ", "")
        score = 0
        if "대학원" in school:
            score += 1
        elif "대학교" in school or school.endswith("대학"):
            score += 2
        if major:
            score += 1
        if status:
            score += 1

        if score > best_edu_score:
            best_edu_score = score
            edu_date = period
            edu_str = _build_edu_str(school, major, status, edu_date)

    # 경력: 경력 구간 내(없으면 전체 스캔) 날짜 범위 + 회사/직무 추출
    career_period = ""
    career_company = ""
    career_job = ""
    end = career_end if career_end > career_start >= 0 else len(lines)
    clines = lines[career_start + 1:end] if career_start >= 0 else lines

    ym_positions: list[tuple[int, str, str]] = []
    for idx, ln in enumerate(clines):
        m = re.fullmatch(r"(\d{4})[.\-]?(\d{1,2})", ln)
        if m:
            ym_positions.append((idx, m.group(1), m.group(2)))

    # 가능한 모든 pair 중에서 직업 문맥이 있는 첫 pair 선택
    for p in range(0, len(ym_positions) - 1, 2):
        _, s_y, s_m = ym_positions[p]
        e_idx, e_y, e_m = ym_positions[p + 1]
        tail = clines[e_idx + 1:e_idx + 12]
        tail_text = " ".join(tail)
        if not any(k in tail_text for k in _CAREER_KWS) and "회사" not in tail_text:
            continue
        career_period = f"{s_y}.{s_m}~{e_y}.{e_m}"
        filtered = [
            x for x in tail
            if x not in {"-", "회사명", "직급", "부서", "담당업무", "연봉", "년", "월"}
            and not _line_has_header(x)
        ]
        if filtered:
            career_company = filtered[0]
        # 직무는 키워드 우선
        career_job = next(
            (
                x for x in filtered[1:]
                if re.search(r"(개발|분석|관리|기획|설계|연구|엔지니어|담당|업무|부서|직무|QA|테스트)", x, re.I)
            ),
            filtered[1] if len(filtered) > 1 else "",
        )
        break

    return edu_str, career_period, career_company, career_job


def _parse_pdf_sections(text: str) -> tuple[str, str, str, str]:
    """PDF 표형 이력서의 학력/경력 보강 추출"""
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
    if not lines:
        return "", "", "", ""

    def _find_index(keyword: str) -> int:
        for i, ln in enumerate(lines):
            if keyword in ln:
                return i
        return -1

    def _is_header(line: str) -> bool:
        return any(h in line for h in _NOISE_LINE_KWS)

    edu_start = _find_index("학력사항")
    career_start = _find_index("경력사항")
    career_end = _find_index("해외경험")

    # 학력
    edu_str = ""
    best_score = -1
    edu_date = ""
    edu_lines = (
        lines[edu_start + 1: career_start if career_start > edu_start else len(lines)]
        if edu_start >= 0 else lines
    )
    for i, ln in enumerate(edu_lines):
        if not re.search(r"(대학교|대학원|대학)$", ln):
            continue
        if _is_header(ln) or ln in _EDU_SKIP:
            continue

        major = ""
        for nxt in edu_lines[i + 1:i + 5]:
            if _is_header(nxt):
                continue
            if any(k in nxt for k in ["학과", "학부", "전공", "계열"]):
                major = nxt
                break

        window = " ".join(edu_lines[max(0, i - 10): i + 2])
        ranges = _DATE_RANGE.findall(window)
        period = ranges[-1] if ranges else ""
        if not period:
            yms = _YEARMONTH.findall(window)
            if len(yms) >= 2:
                period = f"{yms[-2]}~{yms[-1]}"
        if not period:
            continue

        status = next((gs for gs in _GRAD_STATUS if gs in " ".join(edu_lines[i:i + 8])), "")
        school = ln.replace(" ", "")
        score = _score_edu_candidate(school, major, status, period)
        if score > best_score:
            best_score = score
            edu_date = period
            edu_str = _build_edu_str(school, major, status, edu_date)

    # 경력
    career_period = ""
    career_company = ""
    career_job = ""
    end = career_end if career_end > career_start >= 0 else len(lines)
    clines = lines[career_start + 1:end] if career_start >= 0 else lines

    for i, ln in enumerate(clines):
        if not _DATE_RANGE.search(ln):
            continue
        tail = clines[i + 1:i + 12]
        tail_text = " ".join(tail)
        if "회사" not in tail_text and not any(k in tail_text for k in _CAREER_KWS):
            continue
        career_period = _DATE_RANGE.search(ln).group(1)
        filtered = [
            x for x in tail
            if not _is_header(x) and x not in {"-", "/", "`8000"} and not _DATE_RANGE.search(x)
        ]
        if filtered:
            career_company = filtered[0]
        career_job = next(
            (
                x for x in filtered[1:]
                if re.search(r"(개발|분석|관리|기획|설계|연구|엔지니어|담당|업무|부서|직무|QA|테스트)", x, re.I)
            ),
            filtered[1] if len(filtered) > 1 else "",
        )
        break

    return edu_str, career_period, career_company, career_job


def _find_col(header_row: list[str], keywords: list[str]) -> int:
    """헤더 행에서 키워드가 포함된 첫 번째 컬럼 인덱스 반환 (-1 = 없음)"""
    for j, cell in enumerate(header_row):
        if any(kw in cell for kw in keywords):
            return j
    return -1


def _build_edu_str(school: str, major: str, status: str, period: str) -> str:
    parts = [p for p in [school, major, status] if p and p.strip() and p not in _EDU_SKIP]
    if period:
        parts.append(f"({period})")
    return " ".join(parts)


def _score_edu_candidate(school: str, major: str, status: str, period: str) -> int:
    """최종학력 후보 점수: 기간 필수, 학위/정보량 우선"""
    if not period:
        return -1
    score = 0
    if "대학원" in school:
        score += 1
    elif "대학교" in school or school.endswith("대학"):
        score += 2
    if major:
        score += 1
    if status:
        score += 1
    return score


def _is_bad_edu_value(value: str) -> bool:
    if not value.strip():
        return True
    if value.strip().startswith("("):
        return True
    return not bool(re.search(r"(대학교|대학원|대학)", value))


def _finalize_parsed_record(parsed: dict) -> dict:
    """추출 실패/누락 시 안전값으로 보정"""
    required_keys = [
        "이름", "생년월일", "연락처", "이메일", "주소",
        "원문 지원직무", "최종학력", "경력기간", "경력회사", "경력직무",
    ]

    normalized: dict[str, str] = {}
    for k in required_keys:
        v = parsed.get(k, "")
        if v is None:
            v = ""
        normalized[k] = str(v).strip()

    # 카테고리성 정보는 누락 시 '기타'로 보정
    if not normalized["원문 지원직무"]:
        normalized["원문 지원직무"] = "기타"
    if not normalized["최종학력"]:
        normalized["최종학력"] = "기타"
    if not normalized["경력직무"] and normalized["경력기간"]:
        normalized["경력직무"] = "기타"
    if not normalized["경력회사"] and normalized["경력기간"]:
        normalized["경력회사"] = "기타"

    # 부가 필드는 기존 값 유지(추출신뢰도는 사용하지 않음)
    for k, v in parsed.items():
        if k not in normalized and k != "추출신뢰도":
            normalized[k] = "" if v is None else str(v).strip()
    return normalized


def _extract_json_object(text: str) -> dict:
    """모델 응답에서 첫 JSON 객체를 추출"""
    text = text.strip()
    try:
        return json.loads(text)
    except Exception:
        pass
    start = text.find("{")
    end = text.rfind("}")
    if start >= 0 and end > start:
        try:
            return json.loads(text[start:end + 1])
        except Exception:
            return {}
    return {}


def _should_anonymize_ai() -> bool:
    return os.getenv("AI_ANONYMIZE", "true").strip().lower() in {"1", "true", "yes", "on"}


def _mask_sensitive_text(text: str) -> tuple[str, dict[str, str]]:
    """AI 전송 전 민감정보를 토큰으로 치환"""
    masked = text
    token_map: dict[str, str] = {}
    seq = 1

    def _add_token(raw: str, kind: str) -> str:
        nonlocal seq
        raw = raw.strip()
        if not raw:
            return raw
        for tok, original in token_map.items():
            if original == raw:
                return tok
        token = f"__{kind}_{seq}__"
        seq += 1
        token_map[token] = raw
        return token

    # 1) 이메일
    email_pat = re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}")
    for m in sorted(set(email_pat.findall(masked)), key=len, reverse=True):
        masked = masked.replace(m, _add_token(m, "EMAIL"))

    # 2) 전화번호
    phone_pat = re.compile(r"(?:010|011|016|017|018|019|02|0[3-6]\d|070)[-\s]?\d{3,4}[-\s]?\d{4}")
    for m in sorted(set(phone_pat.findall(masked)), key=len, reverse=True):
        masked = masked.replace(m, _add_token(m, "PHONE"))

    # 3) 생년월일/주민번호 유사 패턴
    birth_pat = re.compile(r"\b\d{4}[.\-/]\d{1,2}[.\-/]\d{1,2}\b")
    for m in sorted(set(birth_pat.findall(masked)), key=len, reverse=True):
        masked = masked.replace(m, _add_token(m, "BIRTH"))

    rrn_pat = re.compile(r"\b\d{6}[-\s]?\d{7}\b")
    for m in sorted(set(rrn_pat.findall(masked)), key=len, reverse=True):
        masked = masked.replace(m, _add_token(m, "PID"))

    # 4) 기본정보 기반 이름/주소
    basic = _extract_basic(masked)
    if basic.get("이름"):
        name = basic["이름"].strip()
        if len(name) >= 2:
            masked = masked.replace(name, _add_token(name, "NAME"))
    if basic.get("생년월일"):
        birth = basic["생년월일"].strip()
        if birth:
            masked = masked.replace(birth, _add_token(birth, "BIRTH"))
    if basic.get("주소"):
        addr = basic["주소"].strip()
        if len(addr) >= 4:
            masked = masked.replace(addr, _add_token(addr, "ADDR"))

    return masked, token_map


def _restore_tokens(value: str, token_map: dict[str, str]) -> str:
    restored = value
    for token, original in token_map.items():
        restored = restored.replace(token, original)
    return restored


def _call_openai(prompt: str) -> str:
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("OPENAI_API_KEY가 없습니다.")
    model = os.getenv("OPENAI_MODEL", "").strip()
    if not model:
        raise RuntimeError("OPENAI_MODEL이 없습니다.")
    payload = {
        "model": model,
        "temperature": 0,
        "messages": [
            {"role": "system", "content": "당신은 이력서 정보를 JSON으로 추출하는 엔진입니다."},
            {"role": "user", "content": prompt},
        ],
    }
    req = urllib.request.Request(
        url="https://api.openai.com/v1/chat/completions",
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=60) as resp:
        data = json.loads(resp.read().decode("utf-8", errors="replace"))
    return data["choices"][0]["message"]["content"]


def _call_anthropic(prompt: str) -> str:
    api_key = os.getenv("ANTHROPIC_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("ANTHROPIC_API_KEY가 없습니다.")
    model = os.getenv("ANTHROPIC_MODEL", "").strip()
    if not model:
        raise RuntimeError("ANTHROPIC_MODEL이 없습니다.")
    payload = {
        "model": model,
        "max_tokens": 1200,
        "temperature": 0,
        "messages": [{"role": "user", "content": prompt}],
    }
    req = urllib.request.Request(
        url="https://api.anthropic.com/v1/messages",
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "x-api-key": api_key,
            "anthropic-version": "2023-06-01",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=60) as resp:
        data = json.loads(resp.read().decode("utf-8", errors="replace"))
    blocks = data.get("content", [])
    text_blocks = [b.get("text", "") for b in blocks if b.get("type") == "text"]
    return "\n".join(text_blocks).strip()


def _call_gemini(prompt: str) -> str:
    api_key = os.getenv("GEMINI_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("GEMINI_API_KEY가 없습니다.")
    model = os.getenv("GEMINI_MODEL", "").strip()
    if not model:
        raise RuntimeError("GEMINI_MODEL이 없습니다.")
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={api_key}"
    payload = {
        "generationConfig": {"temperature": 0},
        "contents": [{"parts": [{"text": prompt}]}],
    }
    req = urllib.request.Request(
        url=url,
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=60) as resp:
        data = json.loads(resp.read().decode("utf-8", errors="replace"))
    cands = data.get("candidates", [])
    if not cands:
        return ""
    parts = cands[0].get("content", {}).get("parts", [])
    return "\n".join(p.get("text", "") for p in parts if p.get("text")).strip()


def _should_use_ai() -> bool:
    return os.getenv("USE_AI", "false").strip().lower() in {"1", "true", "yes", "on"}


def _resolve_ai_providers() -> list[str]:
    raw = os.getenv("AI_PROVIDER", "").strip().lower()
    available: list[str] = []
    if os.getenv("OPENAI_API_KEY", "").strip():
        available.append("openai")
    if os.getenv("ANTHROPIC_API_KEY", "").strip():
        available.append("anthropic")
    if os.getenv("GEMINI_API_KEY", "").strip():
        available.append("gemini")

    if not raw:
        return available[:1]
    if raw in {"all", "ensemble", "multi"}:
        return available

    requested = [x.strip() for x in raw.split(",") if x.strip()]
    resolved = [p for p in requested if p in {"openai", "anthropic", "gemini"} and p in available]
    return resolved


def _aggregate_ai_results(results: list[tuple[str, dict]]) -> dict:
    if not results:
        return {}
    fields = [
        "이름", "생년월일", "연락처", "이메일", "주소",
        "원문 지원직무", "최종학력", "경력기간", "경력회사", "경력직무",
    ]
    merged: dict[str, str] = {}
    for field in fields:
        votes: dict[str, int] = {}
        for _, r in results:
            v = str(r.get(field, "")).strip()
            if not v:
                continue
            votes[v] = votes.get(v, 0) + 1
        if not votes:
            merged[field] = ""
            continue
        # 다수결 우선, 동률이면 길이가 긴 값 우선
        best = sorted(votes.items(), key=lambda kv: (kv[1], len(kv[0])), reverse=True)[0][0]
        merged[field] = best

    # 최종학력은 유효성 체크를 통과한 값만 유지
    if merged.get("최종학력") and not _is_valid_final_edu(merged["최종학력"]):
        merged["최종학력"] = ""
    return merged


def _ai_extract_fields(text: str, filename: str, filetype: str) -> dict:
    """AI API로 이력서 필드 추출"""
    providers = _resolve_ai_providers()
    if not providers:
        return {}

    prompt_text = text[:12000]
    token_map: dict[str, str] = {}
    if _should_anonymize_ai():
        prompt_text, token_map = _mask_sensitive_text(prompt_text)

    prompt = (
        "아래 텍스트에서 이력서 정보를 추출해 JSON만 반환하세요.\n"
        "규칙:\n"
        "1) 최종학력은 반드시 학교명+기간이 같이 있는 항목만 인정.\n"
        "2) '단과대학', '학교명' 같은 헤더 문구는 값으로 쓰지 말 것.\n"
        "3) 없는 값은 빈 문자열.\n"
        "4) 키는 정확히 아래만 사용.\n"
        "5) __NAME_n__, __PHONE_n__, __EMAIL_n__, __ADDR_n__, __BIRTH_n__, __PID_n__ 같은 토큰이 보이면 그대로 유지.\n\n"
        "{\n"
        '  "이름": "", "생년월일": "", "연락처": "", "이메일": "", "주소": "",\n'
        '  "원문 지원직무": "", "최종학력": "", "경력기간": "", "경력회사": "", "경력직무": ""\n'
        "}\n\n"
        f"[파일명] {filename}\n"
        f"[파일형식] {filetype}\n"
        "[문서텍스트]\n"
        f"{prompt_text}"
    )

    parsed_results: list[tuple[str, dict]] = []
    for provider in providers:
        try:
            if provider == "openai":
                raw = _call_openai(prompt)
            elif provider == "anthropic":
                raw = _call_anthropic(prompt)
            else:
                raw = _call_gemini(prompt)
            parsed = _extract_json_object(raw)
            if isinstance(parsed, dict):
                parsed_results.append((provider, parsed))
        except Exception as e:
            print(f"[WARN] AI 추출 실패({provider}): {e}")

    if not parsed_results:
        return {}
    allowed = {
        "이름", "생년월일", "연락처", "이메일", "주소",
        "원문 지원직무", "최종학력", "경력기간", "경력회사", "경력직무",
    }
    normalized = []
    for provider, parsed in parsed_results:
        normalized.append((provider, {k: str(parsed.get(k, "")).strip() for k in allowed}))

    result = _aggregate_ai_results(normalized) if len(normalized) > 1 else normalized[0][1]
    if token_map:
        for k, v in list(result.items()):
            result[k] = _restore_tokens(v, token_map)
    return result


def _is_valid_final_edu(value: str) -> bool:
    if _is_bad_edu_value(value):
        return False
    return bool(re.search(r"\(\d{4}[.\-]\d{1,2}\s*[~\-]\s*\d{4}[.\-]\d{1,2}", value))


def _normalize_for_match(text: str) -> str:
    return re.sub(r"\s+", "", text).lower()


def _exists_in_source(candidate: str, source_text: str) -> bool:
    if not candidate.strip():
        return False
    c = _normalize_for_match(candidate)
    s = _normalize_for_match(source_text)
    if c in s:
        return True
    # 기간 표기 차이(하이픈/~/점) 완화
    c2 = c.replace("~", "").replace("-", "").replace(".", "")
    s2 = s.replace("~", "").replace("-", "").replace(".", "")
    return c2 in s2


def _merge_ai_result(base: dict, ai: dict, source_text: str) -> dict:
    if not ai:
        return base
    merged = dict(base)
    for k in ["이름", "생년월일", "연락처", "이메일", "주소", "원문 지원직무"]:
        if not merged.get(k) and ai.get(k) and _exists_in_source(ai.get(k, ""), source_text):
            merged[k] = ai[k]
    if (
        ai.get("최종학력")
        and _exists_in_source(ai.get("최종학력", ""), source_text)
        and (_is_bad_edu_value(merged.get("최종학력", "")) and _is_valid_final_edu(ai["최종학력"]))
    ):
        merged["최종학력"] = ai["최종학력"]
    for k in ["경력기간", "경력회사", "경력직무"]:
        if not merged.get(k) and ai.get(k) and _exists_in_source(ai.get(k, ""), source_text):
            merged[k] = ai[k]
    merged.pop("추출신뢰도", None)
    return merged


# ─── DOCX 직접 파싱 ─────────────────────────────────────────────────────────

def extract_from_docx(filepath: str, original_filename: str) -> dict:
    """python-docx로 DOCX 테이블 구조를 직접 파싱 (LibreOffice 불필요)"""
    try:
        from docx import Document as _Doc
    except ImportError:
        subprocess.run(["pip", "install", "python-docx", "-q"], check=False)
        from docx import Document as _Doc

    doc = _Doc(filepath)

    # ── 1. 모든 단락 & 테이블을 텍스트로 ──
    para_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    all_tables: list[list[list[str]]] = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            # 연속 병합 셀 중복 제거 후 저장
            raw = [c.text.strip().replace("\n", " ").replace("\r", " ") for c in row.cells]
            deduped: list[str] = []
            prev = None
            for cell in raw:
                if cell != prev:
                    deduped.append(cell)
                    prev = cell
            rows.append(deduped)
        all_tables.append(rows)

    # 전체 텍스트 (기본정보 파싱용) - 테이블을 먼저 배치해 구조화된 정보 우선 매칭
    table_flat = "\n".join(
        " ".join(cell for cell in row) for rows in all_tables for row in rows
    )
    full_text = table_flat + "\n" + para_text

    # ── 2. 기본정보: 셀 기반 먼저 추출 → 누락 필드만 텍스트 정규식으로 보완 ──
    # (셀 기반을 먼저 실행해야 병합 셀·사진 칸 오염을 차단할 수 있음)
    info: dict[str, str] = {k: "" for k in
                            ["이름", "생년월일", "연락처", "이메일", "주소", "원문 지원직무"]}

    # 테이블 셀에서 기본 필드 직접 추출 (병합 셀 / 사진 칸 오염 방지)
    _PHOTO_SKIP = {"사진", "사  진", "사 진", "photo", "Photo"}

    def _is_noise_cell(v: str) -> bool:
        if not v.strip():
            return True
        if any(p in v for p in _PHOTO_SKIP):
            return True
        return _norm(v).lower() in (_NAME_NORM | _PHONE_NORM | _ADDR_NORM | _EMAIL_NORM | _BIRTH_NORM)

    def _pick_value(v: str, want_pattern: str | None = None) -> str:
        if _is_noise_cell(v):
            return ""
        if want_pattern:
            m = re.search(want_pattern, v)
            return m.group(1) if m else ""
        return v.strip()

    def _find_value_near_key(
        rows: list[list[str]],
        row_idx: int,
        key_idx: int,
        want_pattern: str | None = None,
    ) -> str:
        """키 셀 주변에서 값을 찾음(오른쪽 우선, 아래/위 폴백)"""
        current = rows[row_idx]

        # 1) 같은 행 오른쪽
        for j in range(key_idx + 1, len(current)):
            picked = _pick_value(current[j], want_pattern)
            if picked:
                return picked

        # 2) 바로 아래/위 동일 컬럼 (양식에 따라 값이 세로 배치되는 경우)
        for offset in (1, -1):
            r = row_idx + offset
            if 0 <= r < len(rows) and key_idx < len(rows[r]):
                picked = _pick_value(rows[r][key_idx], want_pattern)
                if picked:
                    return picked

        # 3) 바로 아래/위 오른쪽 인접 컬럼
        for offset in (1, -1):
            r = row_idx + offset
            if 0 <= r < len(rows):
                for j in range(key_idx + 1, len(rows[r])):
                    picked = _pick_value(rows[r][j], want_pattern)
                    if picked:
                        return picked

        return ""

    def _norm(s: str) -> str:
        return re.sub(r"\s+", "", s)

    _NAME_NORM   = {"성명", "이름", "name"}
    _PHONE_NORM  = {"연락처", "휴대폰", "핸드폰", "mobile", "전화번호", "모바일",
                    "휴대전화", "핸드폰"}
    _ADDR_NORM   = {"주소", "거주지", "address"}
    _EMAIL_NORM  = {"email", "e-mail", "이메일"}
    _BIRTH_NORM  = {"생년월일", "생년", "birthdate"}

    for rows in all_tables:
        for row_idx, row in enumerate(rows):
            for i, cell in enumerate(row):
                nk = _norm(cell).lower()
                if nk in _NAME_NORM and not info["이름"]:
                    info["이름"] = _find_value_near_key(rows, row_idx, i, r"([가-힣]{2,5})")
                elif nk in _PHONE_NORM and not info["연락처"]:
                    _mobile_pat   = r"(?:010|011|016|017|018|019)[-\s]?\d{3,4}[-\s]?\d{4}"
                    _landline_pat = (
                        r"(?:02|031|032|033|041|042|043|044|051|052|053"
                        r"|054|055|061|062|063|064|070)[-\s]?\d{3,4}[-\s]?\d{4}"
                    )
                    # 행 전체에서 휴대폰 번호 먼저 탐색, 없으면 지역번호
                    row_text = " ".join(row)
                    mob_m  = re.search(_mobile_pat,   row_text)
                    land_m = re.search(_landline_pat, row_text)
                    if mob_m:
                        info["연락처"] = mob_m.group()
                    elif land_m:
                        info["연락처"] = land_m.group()
                    else:
                        near_phone = _find_value_near_key(rows, row_idx, i)
                        mob_m = re.search(_mobile_pat, near_phone)
                        land_m = re.search(_landline_pat, near_phone)
                        if mob_m:
                            info["연락처"] = mob_m.group()
                        elif land_m:
                            info["연락처"] = land_m.group()
                elif nk in _ADDR_NORM and not info["주소"]:
                    raw_addr = _find_value_near_key(rows, row_idx, i)
                    # 사진 칸 텍스트 후처리 제거: "사  진 (3cm x 4cm)" 등
                    raw_addr = re.sub(r'\s*사\s*진\s*[\(\（].*?[\)\）]', '', raw_addr).strip()
                    raw_addr = re.sub(r'\s*Photo\s*[\(\（].*?[\)\）]', '', raw_addr, flags=re.IGNORECASE).strip()
                    info["주소"] = raw_addr
                elif nk in _EMAIL_NORM and not info["이메일"]:
                    info["이메일"] = _find_value_near_key(
                        rows, row_idx, i, r"([a-zA-Z0-9._%+\-가-힣]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,})"
                    )
                elif nk in _BIRTH_NORM and not info["생년월일"]:
                    info["생년월일"] = _find_value_near_key(rows, row_idx, i, r"(\d{4}[.\-/]\d{2}[.\-/]\d{2})")

    # 셀 기반으로 못 찾은 필드만 텍스트 정규식으로 보완
    text_info = _extract_basic(full_text)
    for k, v in text_info.items():
        if not info.get(k):
            info[k] = v

    # ── 3. 학력/경력 테이블 찾기 ──
    EDU_HDR_KWS = ["학교명", "재학기간", "학력구분", "전공", "졸업구분", "학력", "기간", "학과명"]
    CAREER_HDR_KWS = ["회사명", "재직기간", "담당", "직무", "경력", "업무", "재직부서"]

    edu_rows: list[list[str]] | None = None
    career_rows: list[list[str]] | None = None

    for rows in all_tables:
        if not rows:
            continue
        hdr = " ".join(rows[0])
        edu_score = sum(1 for kw in EDU_HDR_KWS if kw in hdr)
        career_score = sum(1 for kw in CAREER_HDR_KWS if kw in hdr)
        if edu_score >= 2 and edu_rows is None:
            edu_rows = rows
        if career_score >= 2 and career_rows is None:
            career_rows = rows

    # ── 4. 학력 파싱 ──
    final_edu = ""
    edu_date = ""

    if edu_rows:
        header = edu_rows[0]
        best_score = -1

        school_col = _find_col(header, ["학교명", "학교"])
        major_col  = _find_col(header, ["전공", "학과명"])
        period_col = _find_col(header, ["재학기간", "기간"])
        status_col = _find_col(header, ["졸업구분"])   # '졸업' 단독 컬럼은 날짜값이므로 제외
        level_col  = _find_col(header, ["학력구분", "구분"])

        for row in edu_rows[1:]:
            row_text = " ".join(row)
            if not any(kw in row_text for kw in ["대학교", "대학원", "대학"]):
                continue

            # 학교명
            school = ""
            if school_col >= 0 and school_col < len(row):
                school = row[school_col].strip()
            if not school or school in _EDU_SKIP:
                sm = _SCHOOL.search(row_text)
                school = sm.group(1).replace(" ", "") if sm else ""
            else:
                sm = _SCHOOL.search(school)
                if sm:
                    school = sm.group(1).replace(" ", "")

            # 전공
            major = ""
            if major_col >= 0 and major_col < len(row):
                major = row[major_col].strip()
                if major in _EDU_SKIP:
                    major = ""

            # 기간: period 컬럼 → 날짜 범위 패턴 → 별도 입학/졸업 셀 조합
            period = ""
            if period_col >= 0 and period_col < len(row):
                period = row[period_col].strip()
                if period in _EDU_SKIP:
                    period = ""
            if not period:
                dr = _DATE_RANGE.search(row_text)
                if dr:
                    period = dr.group(1)
            if not period:
                # 일반사무 포맷: 입학 연월 + 졸업 연월이 별도 셀
                yms = _YEARMONTH.findall(row_text)
                if len(yms) >= 2:
                    period = f"{yms[0]}~{yms[-1]}"
                elif yms:
                    period = yms[0]
            candidate_period = period

            # 졸업 상태
            status = ""
            if status_col >= 0 and status_col < len(row):
                status = row[status_col].strip()
                # 날짜처럼 생긴 값(2016.02 등)이나 헤더값은 무시
                if status in _EDU_SKIP or status == "-" or _YEARMONTH.search(status):
                    status = ""
            if not status:
                status = next((gs for gs in _GRAD_STATUS if gs in row_text), "")

            score = _score_edu_candidate(school, major, status, candidate_period)
            if score > best_score:
                best_score = score
                edu_date = candidate_period
                final_edu = _build_edu_str(school, major, status, edu_date)

    # ── 5. 경력 파싱 (기간 + 회사 + 직무) ──
    _CAREER_HDR_VALS = {
        "회사명", "재직기간", "담당", "직무", "경력", "업무", "재직부서",
        "부서", "직위", "직급", "기간", "회사", "기관명", "구분", "활동기간",
    }
    career_entries: list[dict] = []

    if career_rows:
        header = career_rows[0]

        # 기간 컬럼
        period_col = _find_col(header, ["재직기간", "기간", "활동기간"])

        # 회사명 컬럼
        company_col = _find_col(header, ["회사명", "회사", "기관명"])

        # 직무 컬럼 (우선순위: 직무/담당 > 업무 > 부서/재직부서 > 직위/직급)
        job_col = _find_col(header, ["직무", "담당"])
        if job_col < 0:
            job_col = _find_col(header, ["업무"])
        if job_col < 0:
            job_col = _find_col(header, ["부서", "재직부서"])
        if job_col < 0:
            job_col = _find_col(header, ["직위", "직급"])

        for row in career_rows[1:]:
            row_text = " ".join(row)

            # 기간
            period = ""
            if period_col >= 0 and period_col < len(row):
                period = row[period_col].strip()
                if period in _EDU_SKIP or not re.search(r"\d{4}", period):
                    period = ""
            if not period:
                dr = _DATE_RANGE.search(row_text)
                if dr:
                    period = dr.group(1)
            if not period:
                continue  # 기간 없으면 경력 행 아님

            # 회사명
            company = ""
            if company_col >= 0 and company_col < len(row):
                company = row[company_col].strip()
                if company in _CAREER_HDR_VALS:
                    company = ""

            # 직무
            job = ""
            if job_col >= 0 and job_col < len(row):
                job = row[job_col].strip()
                if job in _CAREER_HDR_VALS:
                    job = ""
                if len(job) > 60:
                    job = job[:60] + "…"

            career_entries.append({"period": period, "company": company, "job": job})

    career_period  = "\n".join(e["period"]  for e in career_entries)
    career_company = "\n".join(e["company"] for e in career_entries)
    career_job     = "\n".join(e["job"]     for e in career_entries)

    # ── 6. 테이블에서 못 찾은 경우 텍스트 기반 폴백 ──
    all_dates = _DATE_RANGE.findall(full_text)
    clean_flat = re.sub(r"[ \t]+", " ", full_text)

    if _is_bad_edu_value(final_edu):
        final_edu, edu_date = _parse_edu_from_text(clean_flat, all_dates)

    if not career_period:
        career_period, career_company, career_job = _parse_career_from_text(
            clean_flat, all_dates, edu_date
        )

    # 주소 최종 정제: 사진 칸 텍스트 잔여물 제거 (어떤 경로로 설정됐든 무조건 적용)
    if info.get("주소"):
        info["주소"] = re.sub(r'\s*사\s*진.*$', '', info["주소"]).strip()
        info["주소"] = re.sub(r'\s*Photo.*$', '', info["주소"], flags=re.IGNORECASE).strip()

    parsed = {
        **info,
        "최종학력": final_edu,
        "경력기간": career_period,
        "경력회사": career_company,
        "경력직무": career_job,
        "파일형식": os.path.splitext(original_filename)[1].upper().replace(".", ""),
    }
    if _should_use_ai():
        ai = _ai_extract_fields(clean_flat, original_filename, "DOCX")
        parsed = _merge_ai_result(parsed, ai, clean_flat)
    return _finalize_parsed_record(parsed)


# ─── PDF 파싱 (PyMuPDF) ─────────────────────────────────────────────────────

def extract_from_pdf(doc: fitz.Document, original_filename: str) -> dict:
    """LibreOffice 변환 PDF 또는 원본 PDF 파싱"""
    full_text = ""
    for page in doc:
        blocks = page.get_text("blocks")
        # y좌표 15px 단위로 묶어 같은 행 처리 → x 기준 정렬 (우분투 레이아웃 대응)
        blocks.sort(key=lambda b: (round(b[1] / 15) * 15, b[0]))
        for b in blocks:
            full_text += b[4].replace("\r\n", "\n").replace("\r", "\n") + "\n"

    normalized_text = _normalize_pdf_text(full_text)
    clean_text = re.sub(r"[ \t]+", " ", normalized_text)
    info = _extract_basic(clean_text)
    all_dates = _DATE_RANGE.findall(clean_text)

    final_edu, edu_date = _parse_edu_from_text(clean_text, all_dates)

    career_period, career_company, career_job = _parse_career_from_text(
        clean_text, all_dates, edu_date
    )

    # PDF 표 분해 포맷 보강
    pdf_edu, pdf_career_period, pdf_career_company, pdf_career_job = _parse_pdf_sections(normalized_text)
    if _is_bad_edu_value(final_edu) and pdf_edu:
        final_edu = pdf_edu
    if not career_period and pdf_career_period:
        career_period = pdf_career_period
    if not career_company and pdf_career_company:
        career_company = pdf_career_company
    if not career_job and pdf_career_job:
        career_job = pdf_career_job

    parsed = {
        **info,
        "최종학력": final_edu,
        "경력기간": career_period,
        "경력회사": career_company,
        "경력직무": career_job,
        "파일형식": os.path.splitext(original_filename)[1].upper().replace(".", ""),
    }
    if _should_use_ai():
        ai = _ai_extract_fields(clean_text, original_filename, "PDF")
        parsed = _merge_ai_result(parsed, ai, clean_text)
    return _finalize_parsed_record(parsed)


def extract_from_hwp_pyhwp(filepath: str, original_filename: str) -> dict:
    """pyhwp(hwp5)로 HWP 텍스트를 직접 추출해 파싱"""
    try:
        from hwp5.filestructure import Hwp5File
        from hwp5.recordstream import read_records
    except ImportError:
        subprocess.run(["pip", "install", "pyhwp", "-q"], check=False)
        from hwp5.filestructure import Hwp5File
        from hwp5.recordstream import read_records

    hwp_doc = Hwp5File(filepath)
    text_parts: list[str] = []

    for section_name in hwp_doc.text:
        section_stream = hwp_doc.text[section_name].open()
        for record in read_records(section_stream):
            # 67 = HWPTAG_PARA_TEXT
            if record.get("tagid") != 67:
                continue
            payload = record.get("payload", b"")
            if not payload:
                continue
            decoded = payload.decode("utf-16-le", errors="ignore")
            # HWP 제어문자 제거 (줄바꿈은 유지)
            cleaned = "".join(
                ch if (ch >= " " or ch in "\n\r\t") else " "
                for ch in decoded
            )
            cleaned = re.sub(r"[ \t]+", " ", cleaned).strip()
            if cleaned:
                text_parts.append(cleaned)

    text = "\n".join(text_parts)
    if not text.strip():
        raise ValueError("pyhwp 추출 결과가 비어 있습니다.")

    normalized_text = _normalize_hwp_table_text(text)
    clean_text = re.sub(r"[ \t]+", " ", normalized_text)
    info = _extract_basic(clean_text)
    all_dates = _DATE_RANGE.findall(clean_text)

    final_edu, edu_date = _parse_edu_from_text(clean_text, all_dates)
    career_period, career_company, career_job = _parse_career_from_text(
        clean_text, all_dates, edu_date
    )

    # HWP 표 분해 포맷 보강
    hwp_edu, hwp_career_period, hwp_career_company, hwp_career_job = _parse_hwp_sections(normalized_text)
    if _is_bad_edu_value(final_edu) and hwp_edu:
        final_edu = hwp_edu
    if not career_period and hwp_career_period:
        career_period = hwp_career_period
    if not career_company and hwp_career_company:
        career_company = hwp_career_company
    if not career_job and hwp_career_job:
        career_job = hwp_career_job

    parsed = {
        **info,
        "최종학력": final_edu,
        "경력기간": career_period,
        "경력회사": career_company,
        "경력직무": career_job,
        "파일형식": os.path.splitext(original_filename)[1].upper().replace(".", ""),
    }
    if _should_use_ai():
        ai = _ai_extract_fields(clean_text, original_filename, "HWP")
        parsed = _merge_ai_result(parsed, ai, clean_text)
    return _finalize_parsed_record(parsed)


# ─── 엑셀 저장 ──────────────────────────────────────────────────────────────

_EXCEL_COLUMNS = [
    "이름", "생년월일", "연락처", "이메일", "주소",
    "원문 지원직무", "최종학력",
    "경력기간", "경력회사", "경력직무",
]


# 줄바꿈 wrap_text 적용할 컬럼
_WRAP_COLUMNS = {"경력기간", "경력회사", "경력직무", "최종학력", "주소"}


def _apply_wrap_text(path: str) -> None:
    """저장된 엑셀 파일에 줄바꿈 서식 및 행 높이 자동 조절 적용"""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment

    wb = load_workbook(path)
    ws = wb.active

    # 헤더 행에서 wrap 적용 대상 컬럼 인덱스 파악
    wrap_col_indices: set[int] = set()
    for cell in ws[1]:
        if cell.value in _WRAP_COLUMNS:
            wrap_col_indices.add(cell.column)

    # 데이터 행 전체에 wrap_text 적용
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column in wrap_col_indices:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    # 행 높이: 줄 수에 비례해 자동 설정 (1줄 = 15pt)
    for row in ws.iter_rows(min_row=2):
        max_lines = 1
        for cell in row:
            if cell.column in wrap_col_indices and cell.value:
                lines = str(cell.value).count("\n") + 1
                max_lines = max(max_lines, lines)
        ws.row_dimensions[row[0].row].height = max_lines * 15

    wb.save(path)


def save_to_excel(data_list: list[dict]) -> None:
    df_new = pd.DataFrame(data_list)

    # 누락된 컬럼 빈 문자열로 채우고 지정 순서로 정렬
    for col in _EXCEL_COLUMNS:
        if col not in df_new.columns:
            df_new[col] = ""
    # 이미지 양식 고정: 지정 컬럼만 내보냄(추가 컬럼 제거)
    df_new = df_new[_EXCEL_COLUMNS]

    if not os.path.exists(OUTPUT_PATH):
        df_new.to_excel(OUTPUT_PATH, index=False, engine="openpyxl")
    else:
        df_old = pd.read_excel(OUTPUT_PATH, engine="openpyxl")
        # 기존 파일에도 새 컬럼 추가
        for col in _EXCEL_COLUMNS:
            if col not in df_old.columns:
                df_old[col] = ""
        df_final = (
            pd.concat([df_old, df_new], ignore_index=True)
            .drop_duplicates(subset=["이름", "연락처"], keep="last")
        )
        # 이미지 양식 고정: 지정 컬럼만 유지
        df_final = df_final[_EXCEL_COLUMNS]
        df_final.to_excel(OUTPUT_PATH, index=False, engine="openpyxl")

    _apply_wrap_text(OUTPUT_PATH)


def export_data_list_to_excel_path(data_list: list[dict], path: str) -> None:
    """감지된 항목만 엑셀 파일(path)에 저장(기존 백엔드 `save_to_excel`은 폴더 누적용)."""
    if not data_list:
        pd.DataFrame(columns=_EXCEL_COLUMNS).to_excel(path, index=False, engine="openpyxl")
        return
    df_new = pd.DataFrame(data_list)
    for col in _EXCEL_COLUMNS:
        if col not in df_new.columns:
            df_new[col] = ""
    df_new = df_new[_EXCEL_COLUMNS]
    df_new.to_excel(path, index=False, engine="openpyxl")
    _apply_wrap_text(path)


def parse_resume_file(file_path: str) -> dict:
    """단일 이력서 파일을 파싱해 필드 dict를 반환합니다(API·배치 공통)."""
    original_name = os.path.basename(file_path)
    if original_name == OUTPUT_FILENAME or "temp_conv_" in original_name:
        raise ValueError("이 파일은 처리할 수 없습니다.")
    ext = os.path.splitext(original_name)[1].lower()
    if ext == ".docx":
        return extract_from_docx(file_path, original_name)
    if ext == ".hwp":
        try:
            return extract_from_hwp_pyhwp(file_path, original_name)
        except Exception as e:
            print(f"[WARN] {original_name} pyhwp 직접 파싱 실패, PDF 변환으로 폴백: {e}")
    current_pdf: str | None = None
    is_temporary = False
    if ext == ".pdf":
        current_pdf = file_path
    else:
        temp_name = f"temp_conv_{os.path.splitext(original_name)[0]}.pdf"
        subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "pdf",
             "--outdir", BASE_DIR, file_path],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        gen_pdf = os.path.join(BASE_DIR, os.path.splitext(original_name)[0] + ".pdf")
        current_pdf = os.path.join(BASE_DIR, temp_name)
        if os.path.exists(gen_pdf):
            os.rename(gen_pdf, current_pdf)
            is_temporary = True
    if not (current_pdf and os.path.exists(current_pdf)):
        raise FileNotFoundError("PDF로 변환된 파일을 찾을 수 없습니다.")
    try:
        with fitz.open(current_pdf) as pdf_doc:
            return extract_from_pdf(pdf_doc, original_name)
    finally:
        if is_temporary and current_pdf and os.path.exists(current_pdf):
            try:
                os.remove(current_pdf)
            except OSError:
                pass


# ─── 메인 ───────────────────────────────────────────────────────────────────

def main() -> None:
    target_files: list[str] = []
    for ext in ("*.docx", "*.doc", "*.hwp", "*.pptx", "*.pdf"):
        target_files.extend(glob.glob(os.path.join(BASE_DIR, ext)))
    if not target_files:
        return

    all_parsed: list[dict] = []

    for file_path in target_files:
        original_name = os.path.basename(file_path)
        if original_name == OUTPUT_FILENAME or "temp_conv_" in original_name:
            continue
        try:
            all_parsed.append(parse_resume_file(file_path))
        except Exception as e:
            print(f"[WARN] {original_name} 처리 실패: {e}")

    if all_parsed:
        save_to_excel(all_parsed)
        print(f"완료: {len(all_parsed)}건 추출 → {OUTPUT_FILENAME}")


if __name__ == "__main__":
    main()
