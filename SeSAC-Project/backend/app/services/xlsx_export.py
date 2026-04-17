from __future__ import annotations

import io
from datetime import datetime
from typing import Iterable
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def applications_batch_to_xlsx_bytes(
    *,
    batch_title: str,
    jd_preferred: str,
    rows: Iterable[dict],
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "분류결과"
    bold = Font(bold=True)
    ws.append(["지원서 묶음 이름", batch_title])
    ws.append(["우대·자격 텍스트", jd_preferred])
    ws.append([])
    headers = [
        "파일명",
        "지원 경로",
        "블라인드 등급",
        "우대충족",
        "증거 레벨",
        "직무연관 점수",
        "이력서 완성도",
        "누락 필드",
        "블라인드 요약",
        "우대 근거",
        "우대 가산 패턴 히트",
        "우대 제외 패턴 히트",
        "키워드 태그",
        "인용 스니펫",
        "분석시각",
    ]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        ws.cell(row=4, column=c).font = bold
    missing_fill = PatternFill(fill_type="solid", fgColor="FFF4CCCC")
    for idx, r in enumerate(rows, start=5):
        snippets = r.get("blind_snippets") or []
        if isinstance(snippets, list):
            sn = " | ".join(str(x) for x in snippets[:12])
        else:
            sn = str(snippets)
        flags = r.get("keyword_flags") or []
        if isinstance(flags, list):
            fl = ", ".join(str(x) for x in flags)
        else:
            fl = str(flags)
        analyzed = r.get("analyzed_at")
        if isinstance(analyzed, datetime):
            analyzed_s = analyzed.isoformat()
        else:
            analyzed_s = str(analyzed or "")
        missing_fields = r.get("missing_fields") or []
        if isinstance(missing_fields, list):
            missing_s = ", ".join(str(x) for x in missing_fields)
        else:
            missing_s = str(missing_fields)
        rule_hits = r.get("preferred_rule_hits") or []
        rule_excluded = r.get("preferred_rule_excluded_hits") or []
        ws.append(
            [
                r.get("filename", ""),
                r.get("source_platform", "unknown"),
                r.get("blind_tier", ""),
                "Y" if r.get("preferred_met") else "N",
                r.get("evidence_level", "unknown"),
                float(r.get("role_relevance_score") or 0),
                float(r.get("resume_completeness_score") or 0),
                missing_s,
                r.get("blind_summary", ""),
                r.get("preferred_reason", ""),
                ", ".join(str(x) for x in (rule_hits if isinstance(rule_hits, list) else [])),
                ", ".join(str(x) for x in (rule_excluded if isinstance(rule_excluded, list) else [])),
                fl,
                sn,
                analyzed_s,
            ]
        )
        if missing_s:
            ws.cell(row=idx, column=8).fill = missing_fill
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename_for_batch(batch_id: UUID) -> str:
    return f"applications_{batch_id}.xlsx"


def batch_resume_insights_to_xlsx_bytes(*, batch_title: str, rows: Iterable[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "포지션_이력서_인사이트"
    bold = Font(bold=True)
    headers = [
        "후보자",
        "파일명",
        "AI 추정 지원직무",
        "블라인드 등급",
        "블라인드 판단 이유",
        "우대충족",
        "우대 판단 이유",
        "직무연관도",
        "완성도",
        "누락 필드",
        "중복 의심",
        "중복 사유",
        "지원 경로",
        "증거 레벨",
    ]
    ws.append([f"배치: {batch_title}"])
    ws.append([])
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        ws.cell(row=3, column=c).font = bold
    for r in rows:
        missing = r.get("missing_fields") or []
        if isinstance(missing, list):
            missing_s = ", ".join(str(x) for x in missing)
        else:
            missing_s = str(missing)
        ws.append(
            [
                r.get("candidate_name", ""),
                r.get("filename", ""),
                r.get("inferred_position", ""),
                r.get("blind_tier", ""),
                r.get("blind_summary", ""),
                "Y" if r.get("preferred_met") else "N",
                r.get("preferred_reason", ""),
                float(r.get("role_relevance_score") or 0),
                float(r.get("resume_completeness_score") or 0),
                missing_s,
                "Y" if r.get("duplicate_suspected") else "N",
                r.get("duplicate_reason", ""),
                r.get("source_platform", "unknown"),
                r.get("evidence_level", "unknown"),
            ]
        )
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
